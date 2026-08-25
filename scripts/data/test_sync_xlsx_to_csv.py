#!/usr/bin/env python3
"""Regression tests for the Plan25 XLSX authoring sync."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

import sync_xlsx_to_csv as sync


ROOT = sync.ROOT
DATA = sync.DATA_DIR
SCRIPT = ROOT / "scripts" / "data" / "sync_xlsx_to_csv.py"
CANONICAL = sync.CANONICAL_XLSX
ENEMY_CANONICAL = sync.CANONICAL_ENEMY_XLSX
ENCOUNTER_CANONICAL = sync.CANONICAL_ENCOUNTER_XLSX
AUDIO_CANONICAL = sync.CANONICAL_AUDIO_XLSX
CANONICAL_WORKBOOKS = sync.CANONICAL_WORKBOOKS


class SyncXlsxToCsvTests(unittest.TestCase):
    def run_sync(self, *args: str, env: dict[str, str] | None = None, expect_success: bool = True) -> subprocess.CompletedProcess[str]:
        merged_env = os.environ.copy()
        if env:
            merged_env.update(env)
        result = subprocess.run(
            [sys.executable, str(SCRIPT), *args],
            cwd=ROOT,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            env=merged_env,
        )
        if expect_success and result.returncode != 0:
            self.fail(result.stdout)
        if not expect_success and result.returncode == 0:
            self.fail("command unexpectedly succeeded")
        return result

    def copy_workbook(self, temp_dir: Path) -> Path:
        target = temp_dir / "fixture.xlsx"
        shutil.copy2(CANONICAL, target)
        return target

    def mutate_workbook(self, mutator) -> Path:
        temp_dir = Path(tempfile.mkdtemp(prefix="reecho_xlsx_fixture_"))
        workbook_path = self.copy_workbook(temp_dir)
        wb = load_workbook(workbook_path)
        mutator(wb)
        wb.save(workbook_path)
        self.addCleanup(lambda: shutil.rmtree(temp_dir, ignore_errors=True))
        return workbook_path

    def assert_invalid_workbook(self, mutator, token: str) -> None:
        workbook_path = self.mutate_workbook(mutator)
        result = self.run_sync(
            "--input", str(workbook_path),
            "--input", str(ENEMY_CANONICAL),
            "--input", str(ENCOUNTER_CANONICAL),
            "--input", str(AUDIO_CANONICAL),
            "--check",
            expect_success=False,
        )
        self.assertIn(token, result.stdout)
        self.assertRegex(result.stdout, r":tbl[A-Za-z]+:row \d+:column ")

    def production_bytes(self) -> dict[str, bytes]:
        return {name: (DATA / name).read_bytes() for name in sync.TABLE_TO_CSV.values()}

    def generated_bytes(self) -> dict[str, bytes]:
        with tempfile.TemporaryDirectory(prefix="reecho_xlsx_generated_") as temp:
            _, csv_bytes = sync.generate_package(CANONICAL_WORKBOOKS, Path(temp))
            return csv_bytes

    def make_temp_data_dir(self, initial_bytes: dict[str, bytes] | None = None) -> Path:
        temp_dir = Path(tempfile.mkdtemp(prefix="reecho_xlsx_data_dir_"))
        self.addCleanup(lambda: shutil.rmtree(temp_dir, ignore_errors=True))
        shutil.copy2(DATA / "reecho_data_manifest.csv", temp_dir / "reecho_data_manifest.csv")
        shutil.copy2(DATA / "csv_schema.csv", temp_dir / "csv_schema.csv")
        bytes_by_name = initial_bytes or self.production_bytes()
        for name in sync.TABLE_TO_CSV.values():
            (temp_dir / name).write_bytes(bytes_by_name[name])
        return temp_dir

    def valid_old_bytes(self, changes: dict[str, tuple[bytes, bytes]]) -> dict[str, bytes]:
        old = self.production_bytes()
        for name, (before, after) in changes.items():
            self.assertIn(before, old[name], name)
            old[name] = old[name].replace(before, after, 1)
            self.assertNotEqual(old[name], (DATA / name).read_bytes(), name)
        return old

    def stat_snapshot(self, data_dir: Path) -> dict[str, tuple[int, int]]:
        return {name: ((data_dir / name).stat().st_mtime_ns, (data_dir / name).stat().st_ino) for name in sync.TABLE_TO_CSV.values()}

    def sheet_with_table(self, wb, table_name: str):
        tables = sync.workbook_tables(wb)
        return tables[table_name][0]

    @staticmethod
    def table_cell(wb, table_name: str, row_index: int, column_name: str):
        sheet, table = sync.workbook_tables(wb)[table_name]
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        headers = [sheet.cell(min_row, column).value for column in range(min_col, max_col + 1)]
        return sheet.cell(min_row + 1 + row_index, min_col + headers.index(column_name))

    def set_cell_locked(self, wb, table_name: str, row_index: int, column_name: str, locked: bool) -> None:
        cell = self.table_cell(wb, table_name, row_index, column_name)
        protection = copy(cell.protection)
        protection.locked = locked
        cell.protection = protection

    def remove_validations_for_cell(self, wb, table_name: str, row_index: int, column_name: str) -> None:
        sheet = self.sheet_with_table(wb, table_name)
        address = self.table_cell(wb, table_name, row_index, column_name).coordinate
        sheet.data_validations.dataValidation = [
            validation for validation in sheet.data_validations.dataValidation if address not in validation.cells
        ]

    def test_export_maps_cover_manifest_without_duplicate_ownership(self) -> None:
        workbooks = [load_workbook(path, read_only=False) for path in CANONICAL_WORKBOOKS]
        all_owners = []
        for wb in workbooks:
            owners = sync.read_export_map(wb)
            sync.validate_workbook_protection(wb, owners)
            sync.validate_workbook_data_validations(wb)
            all_owners.extend(owners)
        outputs = [owner.output_csv for owner in all_owners]
        self.assertEqual(sorted(outputs), sorted(sync.TABLE_TO_CSV.values()))
        self.assertEqual(len(outputs), len(set(outputs)))
        main_wb, enemy_wb, encounter_wb, _audio_wb = workbooks
        by_sheet: dict[str, list[str]] = {}
        for owner in all_owners:
            by_sheet.setdefault(owner.sheet, []).append(owner.output_csv)
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(main_wb, "tblCharacters").title]), ["character_aliases.csv", "characters.csv"])
        self.assertEqual(by_sheet[self.sheet_with_table(main_wb, "tblCharacterAbilities").title], ["character_abilities.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(main_wb, "tblCards").title]), ["card_effects.csv", "cards.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(main_wb, "tblWeapons").title]), ["attack_steps.csv", "weapon_types.csv", "weapons.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(main_wb, "tblParts").title]), ["part_effects.csv", "parts.csv", "slot_profiles.csv", "slot_types.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(enemy_wb, "tblEnemies").title]), ["enemies.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(enemy_wb, "tblEnemyAbilities").title]), ["enemy_abilities.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(enemy_wb, "tblBossPhases").title]), ["boss_phases.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(encounter_wb, "tblStages").title]), ["stages.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(encounter_wb, "tblEncounters").title]), ["encounters.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(encounter_wb, "tblEncounterWaves").title]), ["encounter_waves.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(encounter_wb, "tblSpawnProfiles").title]), ["spawn_profiles.csv"])
        self.assertEqual(sorted(by_sheet[self.sheet_with_table(encounter_wb, "tblSpawnPolicy").title]), ["spawn_policy.csv"])
        for wb in workbooks:
            tables = sync.workbook_tables(wb)
            for table_name in set(tables) & sync.AUTHORING_TABLES:
                sheet, table = tables[table_name]
                self.assertTrue(sheet.protection.sheet, table_name)
                self.assertFalse(sheet.protection.insertRows, table_name)
                self.assertFalse(sheet.protection.deleteRows, table_name)
                cells = sheet[table.ref]
                self.assertTrue(all(cell.protection.locked for cell in cells[0]), table_name)
                self.assertTrue(all(not cell.protection.locked for row in cells[1:] for cell in row), table_name)
        for sheet_name in sync.SYSTEM_SHEETS:
            sheet = main_wb[sheet_name]
            self.assertTrue(sheet.protection.sheet, sheet_name)
            self.assertTrue(sheet.protection.insertRows, sheet_name)
            self.assertTrue(sheet.protection.deleteRows, sheet_name)
            self.assertTrue(all(cell.protection.locked for row in sheet.iter_rows() for cell in row), sheet_name)

    def test_output_is_deterministic_and_matches_accepted_csv_bytes(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_xlsx_out_a_") as first, tempfile.TemporaryDirectory(prefix="reecho_xlsx_out_b_") as second:
            self.run_sync("--output-dir", first)
            self.run_sync("--output-dir", second)
            for name in sync.TABLE_TO_CSV.values():
                first_bytes = (Path(first) / name).read_bytes()
                self.assertEqual(first_bytes, (Path(second) / name).read_bytes(), name)
                self.assertEqual(first_bytes, (DATA / name).read_bytes(), name)

    def test_enemy_cross_table_reference_failure_reports_location(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_enemy_fixture_") as temp:
            fixture = Path(temp) / "enemy.xlsx"
            shutil.copy2(ENEMY_CANONICAL, fixture)
            wb = load_workbook(fixture)
            self.table_cell(wb, "tblEnemyAbilities", 0, "OwnerEnemyId").value = "M_Missing"
            wb.save(fixture)
            result = self.run_sync(
                "--input", str(CANONICAL),
                "--input", str(fixture),
                "--input", str(ENCOUNTER_CANONICAL),
                "--input", str(AUDIO_CANONICAL),
                "--check",
                expect_success=False,
            )
            self.assertIn("unknown reference 'M_Missing'", result.stdout)
            self.assertRegex(result.stdout, r"EnemyAbilities:tblEnemyAbilities:row \d+:column OwnerEnemyId")

    def test_enemy_conditional_field_failure_reports_location(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_enemy_fixture_") as temp:
            fixture = Path(temp) / "enemy.xlsx"
            shutil.copy2(ENEMY_CANONICAL, fixture)
            wb = load_workbook(fixture)
            self.table_cell(wb, "tblEnemies", 0, "FuseSeconds").value = 1
            wb.save(fixture)
            result = self.run_sync(
                "--input", str(CANONICAL),
                "--input", str(fixture),
                "--input", str(ENCOUNTER_CANONICAL),
                "--input", str(AUDIO_CANONICAL),
                "--check",
                expect_success=False,
            )
            self.assertIn("non-Bomber fields must use explicit zero", result.stdout)
            self.assertRegex(result.stdout, r"Enemies:tblEnemies:row \d+:column TriggerRadiusCm")

    def test_enemy_shard_drop_contract_rejects_half_or_inverted_ranges(self) -> None:
        self.assert_invalid_workbook(
            lambda wb: setattr(self.table_cell(wb, "tblEnemyShardDrops", 0, "EliteMin"), "value", 10),
            "EliteMin and EliteMax must both be blank or both configured",
        )
        self.assert_invalid_workbook(
            lambda wb: setattr(self.table_cell(wb, "tblEnemyShardDrops", 2, "MeleeMin"), "value", 4),
            "minimum cannot exceed maximum",
        )

    def test_enemy_shard_drop_contract_rejects_duplicate_encounter_rows(self) -> None:
        self.assert_invalid_workbook(
            lambda wb: setattr(self.table_cell(wb, "tblEnemyShardDrops", 1, "EncounterIndex"), "value", 1),
            "duplicate EncounterIndex",
        )

    def test_legacy_enemy_reward_column_is_removed(self) -> None:
        workbook = load_workbook(ENEMY_CANONICAL)
        sheet, table = sync.workbook_tables(workbook)["tblEnemies"]
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        headers = [sheet.cell(min_row, column).value for column in range(min_col, max_col + 1)]
        self.assertNotIn("Reward", headers)

    def test_encounter_anchor_ratio_failure_reports_location(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_encounter_fixture_") as temp:
            fixture = Path(temp) / "encounter.xlsx"
            shutil.copy2(ENCOUNTER_CANONICAL, fixture)
            wb = load_workbook(fixture)
            self.table_cell(wb, "tblEncounters", 2, "PlayerAnchorRatio").value = 0.2
            wb.save(fixture)
            result = self.run_sync(
                "--input", str(CANONICAL),
                "--input", str(ENEMY_CANONICAL),
                "--input", str(fixture),
                "--input", str(AUDIO_CANONICAL),
                "--check",
                expect_success=False,
            )
            self.assertIn("anchor ratios must sum to 1", result.stdout)
            self.assertRegex(result.stdout, r"Encounters:tblEncounters:row \d+:column EchoAnchorRatio")

    def test_encounter_wave_timing_failure_reports_location(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_encounter_fixture_") as temp:
            fixture = Path(temp) / "encounter.xlsx"
            shutil.copy2(ENCOUNTER_CANONICAL, fixture)
            wb = load_workbook(fixture)
            self.table_cell(wb, "tblEncounterWaves", 1, "TriggerSeconds").value = 11
            wb.save(fixture)
            result = self.run_sync(
                "--input", str(CANONICAL),
                "--input", str(ENEMY_CANONICAL),
                "--input", str(fixture),
                "--input", str(AUDIO_CANONICAL),
                "--check",
                expect_success=False,
            )
            self.assertIn("require waves 1..3 at 0/10/20 seconds", result.stdout)
            self.assertRegex(result.stdout, r"EncounterWaves:tblEncounterWaves:row \d+:column WaveIndex")

    def test_check_is_read_only_and_reports_drift(self) -> None:
        generated = self.generated_bytes()
        old = self.valid_old_bytes({"runtime_smoke.csv": (b"42.5", b"43.5")})
        temp_data = self.make_temp_data_dir(old)
        before_stats = self.stat_snapshot(temp_data)
        drift = sync.diff_against_content(generated, temp_data)
        self.assertEqual(drift, ["runtime_smoke.csv"])
        self.assertEqual(old["runtime_smoke.csv"], (temp_data / "runtime_smoke.csv").read_bytes())
        self.assertEqual(before_stats, self.stat_snapshot(temp_data))

    def test_sheet_publish_validates_global_snapshot_and_writes_group(self) -> None:
        generated = self.generated_bytes()
        old = self.valid_old_bytes(
            {
                "weapon_types.csv": (b"Pattern.LongSwordCombo", b"Pattern.LongSwordDashOnly"),
                "weapons.csv": (b"W_J_04,Scythe", b"W_J_04,ScytheOld"),
                "attack_steps.csv": (b"0.5,0.60,150", b"0.55,0.60,150"),
            }
        )
        temp_data = self.make_temp_data_dir(old)
        before_stats = self.stat_snapshot(temp_data)
        selected = ["weapon_types.csv", "weapons.csv", "attack_steps.csv"]
        sync.publish(generated, selected, data_dir=temp_data, final_validator=lambda: None)
        for name in selected:
            self.assertEqual(generated[name], (temp_data / name).read_bytes(), name)
        after_stats = self.stat_snapshot(temp_data)
        for name in set(sync.TABLE_TO_CSV.values()) - set(selected):
            self.assertEqual(old[name], (temp_data / name).read_bytes(), name)
            self.assertEqual(before_stats[name], after_stats[name], name)

    def test_noncanonical_input_cannot_write_production(self) -> None:
        with tempfile.TemporaryDirectory(prefix="reecho_xlsx_noncanonical_") as temp:
            workbook_path = self.copy_workbook(Path(temp))
            result = self.run_sync("--input", str(workbook_path), expect_success=False)
            self.assertIn("Non-canonical --input cannot publish", result.stdout)

    def test_invalid_workbook_cases_fail_with_location(self) -> None:
        self.assert_invalid_workbook(lambda wb: self.sheet_with_table(wb, "tblCharacters").tables.pop("tblCharacters"), "Workbook table is missing")
        self.assert_invalid_workbook(lambda wb: setattr(self.table_cell(wb, "tblCharacters", -1, "HpMax"), "value", "HpMaxBroken"), "Columns do not match")
        self.assert_invalid_workbook(lambda wb: setattr(self.table_cell(wb, "tblCharacters", 0, "HpMax"), "value", "not-a-number"), "finite number")
        self.assert_invalid_workbook(lambda wb: setattr(self.table_cell(wb, "tblCharacters", 0, "DefaultWeaponId"), "value", "W_UNKNOWN"), "DefaultWeaponId")
        self.assert_invalid_workbook(lambda wb: setattr(self.table_cell(wb, "tblCharacterAbilities", 0, "BehaviorId"), "value", "Unknown.Handler"), "behavior id")
        self.assert_invalid_workbook(lambda wb: setattr(self.sheet_with_table(wb, "tblRuntimeSmoke")["D4"], "value", "=1+1"), "Formula cells are not allowed")
        self.assert_invalid_workbook(lambda wb: setattr(self.sheet_with_table(wb, "tblExportMap")["D2"], "value", "../characters.csv"), "plain manifest filename")
        self.assert_invalid_workbook(lambda wb: self.set_cell_locked(wb, "tblCharacters", 0, "RoleId", True), "must be unlocked for authoring")
        self.assert_invalid_workbook(lambda wb: self.set_cell_locked(wb, "tblRuntimeSmoke", 0, "Id", False), "must remain locked")
        self.assert_invalid_workbook(
            lambda wb: self.remove_validations_for_cell(wb, "tblCharacterAbilities", 0, "BehaviorId"),
            "must use an in-cell list validation",
        )

    def test_publish_failure_rolls_back_all_changed_bytes(self) -> None:
        generated = self.generated_bytes()
        old = self.valid_old_bytes({"runtime_smoke.csv": (b"42.5", b"43.5")})
        temp_data = self.make_temp_data_dir(old)
        try:
            os.environ["REECHO_XLSX_FAIL_AFTER_REPLACE"] = "1"
            with self.assertRaises(OSError):
                sync.publish(generated, ["runtime_smoke.csv", "runtime_smoke_effects.csv"], data_dir=temp_data, final_validator=lambda: None)
        finally:
            os.environ.pop("REECHO_XLSX_FAIL_AFTER_REPLACE", None)
        for name, data in old.items():
            self.assertEqual(data, (temp_data / name).read_bytes(), name)
        self.assertFalse((temp_data / sync.TRANSACTION_FILE).exists())

    def test_publish_failure_removes_newly_created_output(self) -> None:
        generated = self.generated_bytes()
        temp_data = self.make_temp_data_dir()
        new_output = temp_data / "stages.csv"
        new_output.unlink()
        try:
            os.environ["REECHO_XLSX_FAIL_AFTER_REPLACE"] = "1"
            with self.assertRaises(OSError):
                sync.publish(generated, ["stages.csv"], data_dir=temp_data, final_validator=lambda: None)
        finally:
            os.environ.pop("REECHO_XLSX_FAIL_AFTER_REPLACE", None)
        self.assertFalse(new_output.exists())
        self.assertFalse((temp_data / sync.TRANSACTION_FILE).exists())

    def test_final_project_validator_failure_rolls_back_all_changed_bytes(self) -> None:
        generated = self.generated_bytes()
        old = self.valid_old_bytes({"runtime_smoke.csv": (b"42.5", b"43.5")})
        temp_data = self.make_temp_data_dir(old)

        def fail_final_validator() -> None:
            raise sync.SyncError("Injected final project validator failure")

        with self.assertRaises(sync.SyncError):
            sync.publish(generated, list(sync.TABLE_TO_CSV.values()), data_dir=temp_data, final_validator=fail_final_validator)
        for name, data in old.items():
            self.assertEqual(data, (temp_data / name).read_bytes(), name)
        self.assertFalse((temp_data / sync.TRANSACTION_FILE).exists())

    def test_legacy_transaction_is_recovered_without_regeneration(self) -> None:
        old = self.valid_old_bytes({"runtime_smoke.csv": (b"42.5", b"43.5")})
        temp_data = self.make_temp_data_dir(old)
        backup_dir = temp_data / ".reecho_csv_publish_backup_test"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / "runtime_smoke.csv"
        backup.write_bytes(old["runtime_smoke.csv"])
        (temp_data / "runtime_smoke.csv").write_bytes(b"corrupted\r\n")
        marker = temp_data / sync.TRANSACTION_FILE
        marker.write_text(json.dumps({"backup_dir": backup_dir.name, "files": ["runtime_smoke.csv"]}), encoding="utf-8")
        sync.recover_transaction(temp_data)
        self.assertEqual(old["runtime_smoke.csv"], (temp_data / "runtime_smoke.csv").read_bytes())
        self.assertFalse(marker.exists())
        self.assertFalse(backup_dir.exists())

    def test_malicious_or_stale_transaction_markers_are_rejected(self) -> None:
        temp_data = self.make_temp_data_dir()
        cases = [
            {"backup_dir": ".reecho_csv_publish_backup_test", "files": ["../runtime_smoke.csv"]},
            {"backup_dir": ".reecho_csv_publish_backup_test", "files": [str(temp_data / "runtime_smoke.csv")]},
            {"backup_dir": ".reecho_csv_publish_backup_test", "files": ["unknown.csv"]},
            {"backup_dir": str(temp_data / ".reecho_csv_publish_backup_test"), "files": ["runtime_smoke.csv"]},
            {"backup_dir": "../.reecho_csv_publish_backup_test", "files": ["runtime_smoke.csv"]},
            {
                "backup_dir": ".reecho_csv_publish_backup_test",
                "files": [{"target": str(temp_data / "runtime_smoke.csv"), "backup": str(temp_data / "x")}],
            },
        ]
        marker = temp_data / sync.TRANSACTION_FILE
        for state in cases:
            marker.write_text(json.dumps(state), encoding="utf-8")
            with self.assertRaises(sync.SyncError):
                sync.recover_transaction(temp_data)


if __name__ == "__main__":
    unittest.main(verbosity=2)
