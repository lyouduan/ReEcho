#!/usr/bin/env python3
"""Author the standalone Plan34 audio catalog workbook.

This script only writes Design/Data/ReEchoAudioEvents.xlsx. It never touches
ReEchoData.xlsx or ReEchoEnemyData.xlsx.
"""
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "Design" / "Data" / "ReEchoAudioEvents.xlsx"
CATALOG_CSV = ROOT / "Content" / "Data" / "audio_events.csv"
COLUMNS = [
    "EventId", "VariantId", "AssetPath", "Bus", "EventType", "Spatial3D", "BaseVolume",
    "PitchMin", "PitchMax", "CooldownSeconds", "MaxConcurrency", "Priority",
    "PausePolicy", "AttenuationMin", "AttenuationMax",
]


def add_table(ws, name: str, headers: list[str], rows: list[list[object]], style: str) -> None:
    for column, header in enumerate(headers, 1):
        cell = ws.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF4F81BD")
        cell.protection = Protection(locked=True)
    for row_index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            ws.cell(row_index, column, value).protection = Protection(locked=False)
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableColumns = [TableColumn(id=index + 1, name=header) for index, header in enumerate(headers)]
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showFirstColumn=False,
                                          showLastColumn=False, showColumnStripes=False)
    ws.add_table(table)


def main() -> int:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    export = wb.create_sheet("_ExportMap")
    add_table(export, "tblExportMap", ["SheetName", "TableName", "AdapterVersion", "OutputCsv", "SortKey"],
              [["AudioEvents", "tblAudioEvents", 1, "audio_events.csv", 30]], "TableStyleMedium9")
    export.protection.sheet = True
    export.protection.insertRows = True
    export.protection.deleteRows = True
    for row in export.iter_rows():
        for cell in row:
            cell.protection = Protection(locked=True)

    with CATALOG_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        catalog_rows = list(csv.DictReader(handle))
    if list(catalog_rows[0]) != COLUMNS:
        raise RuntimeError(f"{CATALOG_CSV} does not match the locked Plan114 15-column schema")
    rows: list[list[object]] = [[row[column] for column in COLUMNS] for row in catalog_rows]

    audio = wb.create_sheet("AudioEvents")
    add_table(audio, "tblAudioEvents", COLUMNS, rows, "TableStyleMedium2")
    audio.protection.sheet = True
    audio.protection.insertRows = False
    audio.protection.deleteRows = False
    validation_values = {
        "Bus": "Music,Ambience,CombatSfx,UiSfx",
        "EventType": "OneShot,Loop",
        "Spatial3D": "true,false",
        "PausePolicy": "PauseWithGame,ContinueOnPause",
    }
    for column_name, values in validation_values.items():
        column = COLUMNS.index(column_name) + 1
        validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=False)
        validation.errorStyle = "stop"
        validation.showErrorMessage = True
        validation.showDropDown = False
        audio.add_data_validation(validation)
        validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{len(rows) + 1}")
    for column, width in enumerate([24, 16, 62, 14, 12, 10, 12, 10, 10, 18, 16, 10, 20, 16, 16], 1):
        audio.column_dimensions[get_column_letter(column)].width = width

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)
    print(f"authored standalone audio workbook: {WORKBOOK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
