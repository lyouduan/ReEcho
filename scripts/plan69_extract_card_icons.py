#!/usr/bin/env python3
"""Plan 69 卡牌 icon 抽取脚本（可复现）。

从《【开普勒】回响数值与构筑体系.xlsx》的 构筑体系G 表 drawing 锚点抽取每张卡的小图标，
按 **卡名称** 与 cards.csv 的 DisplayName 匹配（设计表与游戏数据 Id/名称已漂移，不能按 Id 对应），
落到工作树 Content/SourceArt/UI/Cards/Icon/T_UI_CardIcon_{GAME_CARD_ID}.png。

运行：python scripts/plan69_extract_card_icons.py
前置：需本机存在上述 xlsx（内嵌图片在 xl/drawings/media/）。
"""
import os, zipfile, re, hashlib, csv
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = r"c:\Users\gavynqiu\Documents\miniGame\【开普勒】回响数值与构筑体系.xlsx"
CARDS = os.path.join(ROOT, "Content", "Data", "cards.csv")
WORKTREE_ICON = os.path.join(ROOT, "Content", "SourceArt", "UI", "Cards", "Icon")


def norm(s):
    return re.sub(r"\s+", "", str(s or "").strip())


z = zipfile.ZipFile(XLSX)
wb = ET.fromstring(z.read("xl/workbook.xml"))
sheets = {s.get("name"): s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
          for s in wb.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet")}
wrels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read("xl/_rels/workbook.xml.rels").decode()))
gname = [s for s in sheets if "构筑体系G" in s][0]
sf = "xl/" + wrels[sheets[gname]].lstrip("/").replace("xl/../", "")
sx = z.read(sf).decode()
drid = re.search(r'<drawing r:id="([^"]+)"', sx).group(1)
srel = "xl/worksheets/_rels/" + os.path.basename(sf) + ".rels"
_dr = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(srel).decode()))
_dfile = _dr[drid]
if not _dfile.startswith("xl/"):
    _dfile = "xl/drawings/" + _dfile.split("/")[-1]
dxml = z.read(_dfile).decode("utf-8", "ignore")
_drels2 = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read(_dfile.replace("drawings/", "drawings/_rels/").replace(".xml", ".xml.rels")).decode()))

anchors = re.findall(r"<xdr:(?:oneCell|twoCell)Anchor>.*?</xdr:(?:oneCell|twoCell)Anchor>", dxml, re.S)
row_blip = {}
for anc in anchors:
    frm = re.search(r"<xdr:from>.*?</xdr:from>", anc, re.S)
    if not frm:
        continue
    col = int(re.search(r"<xdr:col>(\d+)</xdr:col>", frm.group(0)).group(1))
    row = int(re.search(r"<xdr:row>(\d+)</xdr:row>", frm.group(0)).group(1))
    blip = re.search(r'r:embed="([^"]+)"', anc)
    if blip and col == 3:
        row_blip[row] = blip.group(1)

wb2 = load_workbook(XLSX, data_only=True)
ws = wb2[gname]
doc_icons = {}
for drow, blip in row_blip.items():
    srow = drow + 1
    name = ws.cell(row=srow, column=3).value
    media = _drels2[blip]
    if not media.startswith("xl/"):
        media = "xl/drawings/media/" + media.split("/")[-1]
    doc_icons[norm(name)] = z.read(media)

game = list(csv.DictReader(open(CARDS, encoding="utf-8-sig")))
name_to_gameid = {norm(r["DisplayName"]): r["Id"] for r in game}

# --- Plan69 diagnostic: which document names carry art, and status of targets ---
import json as _json
_diag = {
    "doc_icon_names": sorted(doc_icons.keys()),
    "target_潮汐回响_in_doc": "潮汐回响" in doc_icons,
    "target_森林回响_in_doc": "森林回响" in doc_icons,
}
with open(os.path.join(os.environ.get("TEMP", "."), "icon_diag.json"), "w", encoding="utf-8") as _f:
    _json.dump(_diag, _f, ensure_ascii=False, indent=2)

os.makedirs(WORKTREE_ICON, exist_ok=True)
used = set()
for dname, data in doc_icons.items():
    gid = name_to_gameid.get(dname)
    if gid and gid not in used:
        with open(os.path.join(WORKTREE_ICON, f"T_UI_CardIcon_{gid}.png"), "wb") as f:
            f.write(data)
        used.add(gid)
missing = [r["Id"] for r in game if r["Id"] not in used and not r["Id"].startswith("FORGE")]
print(f"extracted {len(used)} icons; missing (generic fallback): {missing}")
print(f"FORGE (generic fallback): {[r['Id'] for r in game if r['Id'].startswith('FORGE')]}")
